# scripts/gen_demo_answer_docx.py
# 生成一份规范的学生答卷 Word（面试演示用），内容与 seed_exam_math.py 的演示卷对齐。
# 6 题：客观题 2 对 2 错 + 主观题 1 部分对 1 错，触发完整批改演示（含教师复核）。
# 产出：samples/示例答卷_六年级数学.docx
# 运行：PYTHONPATH=. PYTHONUTF8=1 .venv/Scripts/python.exe scripts/gen_demo_answer_docx.py
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = Path(__file__).parent.parent / "samples" / "示例答卷_六年级数学.docx"

# 题目行（与 seed_exam_math.py QUESTIONS 的 content 第一行一致，不含选项）
QUESTIONS = [
    "一个圆的半径是 3 厘米，它的周长是多少厘米？（π 取 3.14）",
    "计算 3/4 ÷ 2/3 的结果是（ ）",
    "半径为 2 厘米的圆的面积是 12.56 平方厘米。（判断）",
    "两个分数相乘，积一定大于其中任意一个分数。（判断）",
    "一个圆的直径是 10 厘米，请计算它的周长和面积。（π 取 3.14）",
    "把一根绳子剪成两段，第一段长 3/5 米，第二段比第一段短 1/5 米，两根绳子一共长多少米？",
]

# 学生作答：客观题单行；主观题第一行为空（"答："后换行），后续为作答内容
ANSWERS = [
    ["B"],
    ["C"],
    ["正确"],
    ["正确"],
    ["", "周长 = π × d = 3.14 × 10 = 31.4 厘米"],
    ["", "两根绳子一共长 3/5 + 1/5 = 4/5 米"],
]


def _set_font(run, size, bold=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_para(doc, text, size=10.5, bold=False, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    _set_font(p.add_run(text), size, bold)
    return p


def main():
    doc = Document()

    add_para(doc, "数 学 答 卷", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "小学数学 · 圆的周长与面积（演示卷）", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "姓名：张小明        学号：2026001        班级：六年级 4 班", size=10.5)
    add_para(doc, "—" * 40, size=10.5)

    for i, (q, ans) in enumerate(zip(QUESTIONS, ANSWERS), 1):
        add_para(doc, f"第{i}题 {q}", size=10.5, bold=True)
        if ans[0]:
            add_para(doc, "答：" + ans[0], size=10.5)
        else:
            add_para(doc, "答：", size=10.5)
            for ln in ans[1:]:
                add_para(doc, ln, size=10.5)
        doc.add_paragraph()

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"✅ 示例答卷已生成：{OUT}")


if __name__ == "__main__":
    main()
