# scripts/seed_exam_math.py
# 执行：python scripts/seed_exam_math.py
# 用途：生成一份「小学数学演示卷」（贴合项目「K12 AI 教学助教」定位），用于演示试卷批改全流程。
#       一次性完成两件事：
#         ① 幂等灌入 exams + questions + scoring_points（客观 4 题 + 主观 2 题，满分 50）
#         ② 生成学生答卷 docx（含 2 对 2 错客观题 + 1 部分对 1 错主观题，触发教师复核演示）
# 幂等：先清掉该卷的旧提交/批改/题目/得分点，再重灌。

import asyncio
import os

from dotenv import load_dotenv
import asyncpg
from docx import Document

load_dotenv(".env.local")  # 在项目根目录运行本脚本

DB_DSN = (
    f"postgresql://{os.getenv('DB_USER')}:{os.getenv('DB_PASSWORD')}"
    f"@{os.getenv('DB_HOST', 'localhost')}:{os.getenv('DB_PORT', 5433)}"
    f"/{os.getenv('DB_NAME', 'eduagent')}"
)

EXAM_ID = "b3a2c1d0-0000-4000-8000-000000000001"
TENANT_ID = "tenant_default"
DOCX_PATH = "scripts/manual_tests/math_student_answer.docx"

# 6 道题：单选×2 + 判断×2（客观轨）+ 简答×2（主观轨，含得分点）
QUESTIONS = [
    {   # 第1题：单选（圆的周长，答对）
        "id":            "c3a2c1d0-0000-4000-8000-000000000001",
        "question_no":   1,
        "question_type": "single_choice",
        "content":       "一个圆的半径是 3 厘米，它的周长是多少厘米？（π 取 3.14）\n"
                         "A. 9.42 厘米\n"
                         "B. 18.84 厘米\n"
                         "C. 28.26 厘米\n"
                         "D. 12.56 厘米",
        "correct_answer": "B",
        "score":         5,
        "knowledge_tag": "圆的周长",
    },
    {   # 第2题：单选（分数除法，答错）
        "id":            "c3a2c1d0-0000-4000-8000-000000000002",
        "question_no":   2,
        "question_type": "single_choice",
        "content":       "计算 3/4 ÷ 2/3 的结果是（ ）\n"
                         "A. 1/2\n"
                         "B. 9/8\n"
                         "C. 8/9\n"
                         "D. 2/3",
        "correct_answer": "B",
        "score":         5,
        "knowledge_tag": "分数除法",
    },
    {   # 第3题：判断（圆的面积，答对）
        "id":            "c3a2c1d0-0000-4000-8000-000000000003",
        "question_no":   3,
        "question_type": "judge",
        "content":       "半径为 2 厘米的圆的面积是 12.56 平方厘米。（判断）",
        "correct_answer": "正确",
        "score":         5,
        "knowledge_tag": "圆的面积",
    },
    {   # 第4题：判断（分数乘法，答错）
        "id":            "c3a2c1d0-0000-4000-8000-000000000004",
        "question_no":   4,
        "question_type": "judge",
        "content":       "两个分数相乘，积一定大于其中任意一个分数。（判断）",
        "correct_answer": "错误",
        "score":         5,
        "knowledge_tag": "分数乘法",
    },
    {   # 第5题：简答（圆的周长与面积，学生只答了周长）
        "id":            "c3a2c1d0-0000-4000-8000-000000000005",
        "question_no":   5,
        "question_type": "short_answer",
        "content":       "一个圆的直径是 10 厘米，请计算它的周长和面积。（π 取 3.14）",
        "correct_answer": "周长 = πd = 3.14×10 = 31.4 厘米；面积 = πr² = 3.14×5² = 78.5 平方厘米。",
        "score":         15,
        "knowledge_tag": "圆的周长与面积",
    },
    {   # 第6题：简答（分数应用题，学生答错）
        "id":            "c3a2c1d0-0000-4000-8000-000000000006",
        "question_no":   6,
        "question_type": "short_answer",
        "content":       "把一根绳子剪成两段，第一段长 3/5 米，第二段比第一段短 1/5 米，两根绳子一共长多少米？",
        "correct_answer": "第二段长 3/5 - 1/5 = 2/5 米，两根一共 3/5 + 2/5 = 1 米。",
        "score":         15,
        "knowledge_tag": "分数应用题",
    },
]

# 主观题得分点（第5题 5+10、第6题 5+10）
SCORING_POINTS = [
    {"id": "d3a2c1d0-0000-4000-8000-000000000001",
     "question_id": "c3a2c1d0-0000-4000-8000-000000000005",
     "point_desc": "正确计算周长：πd = 3.14×10 = 31.4 厘米", "point_score": 5},
    {"id": "d3a2c1d0-0000-4000-8000-000000000002",
     "question_id": "c3a2c1d0-0000-4000-8000-000000000005",
     "point_desc": "正确计算面积：πr² = 3.14×25 = 78.5 平方厘米", "point_score": 10},
    {"id": "d3a2c1d0-0000-4000-8000-000000000003",
     "question_id": "c3a2c1d0-0000-4000-8000-000000000006",
     "point_desc": "正确求第二段长度：3/5 - 1/5 = 2/5 米", "point_score": 5},
    {"id": "d3a2c1d0-0000-4000-8000-000000000004",
     "question_id": "c3a2c1d0-0000-4000-8000-000000000006",
     "point_desc": "正确求总长：3/5 + 2/5 = 1 米", "point_score": 10},
]


async def seed():
    conn = await asyncpg.connect(DB_DSN)
    print("✅ 数据库连接成功，开始灌入小学数学演示卷...")
    try:
        teacher_id = await conn.fetchval(
            "SELECT id FROM users WHERE username = 'teacher01' LIMIT 1"
        )
        if not teacher_id:
            raise RuntimeError("找不到 teacher01 账号，请先运行 python scripts/seed_data.py")

        async with conn.transaction():
            # ── 幂等清理 ──
            await conn.execute(
                "DELETE FROM exam_reviews WHERE submission_id IN "
                "(SELECT id FROM exam_submissions WHERE exam_id = $1)", EXAM_ID)
            await conn.execute("DELETE FROM exam_submissions WHERE exam_id = $1", EXAM_ID)
            await conn.execute(
                "DELETE FROM scoring_points WHERE question_id IN "
                "(SELECT id FROM questions WHERE exam_id = $1)", EXAM_ID)
            await conn.execute("DELETE FROM questions WHERE exam_id = $1", EXAM_ID)
            await conn.execute("DELETE FROM exams WHERE id = $1", EXAM_ID)

            # ── 灌 exam ──
            await conn.execute(
                "INSERT INTO exams (id, tenant_id, title, description, subject, created_by, is_active) "
                "VALUES ($1, $2, $3, $4, $5, $6, TRUE)",
                EXAM_ID, TENANT_ID,
                "小学数学·圆的周长与面积（演示卷）",
                "演示用：单选×2 + 判断×2 + 简答×2，覆盖客观+主观批改，满分 50",
                "数学",
                teacher_id,
            )

            # ── 灌题 ──
            for q in QUESTIONS:
                await conn.execute(
                    "INSERT INTO questions "
                    "(id, tenant_id, exam_id, question_no, question_type, "
                    " content, correct_answer, score, knowledge_tag) "
                    "VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9)",
                    q["id"], TENANT_ID, EXAM_ID, q["question_no"], q["question_type"],
                    q["content"], q["correct_answer"], q["score"], q["knowledge_tag"],
                )

            # ── 灌得分点 ──
            for sp in SCORING_POINTS:
                await conn.execute(
                    "INSERT INTO scoring_points (id, question_id, point_desc, point_score, is_active) "
                    "VALUES ($1, $2, $3, $4, TRUE)",
                    sp["id"], sp["question_id"], sp["point_desc"], sp["point_score"],
                )

        print("✅ 灌入完成：")
        print(f"   exam:    {EXAM_ID}「小学数学·圆的周长与面积（演示卷）」")
        print(f"   题目数:  {len(QUESTIONS)}（单选2/判断2/简答2，满分 50）")
        print(f"   得分点:  {len(SCORING_POINTS)} 条（简答题）")
    finally:
        await conn.close()


def gen_docx():
    """生成学生答卷 docx：客观题 2 对 2 错，主观题 1 部分对 1 错。"""
    doc = Document()

    doc.add_paragraph("第1题 一个圆的半径是 3 厘米，它的周长是多少厘米？（π 取 3.14）")
    doc.add_paragraph("答：B")
    doc.add_paragraph("")

    doc.add_paragraph("第2题 计算 3/4 ÷ 2/3 的结果是（ ）")
    doc.add_paragraph("答：C")
    doc.add_paragraph("")

    doc.add_paragraph("第3题 半径为 2 厘米的圆的面积是 12.56 平方厘米。（判断）")
    doc.add_paragraph("答：正确")
    doc.add_paragraph("")

    doc.add_paragraph("第4题 两个分数相乘，积一定大于其中任意一个分数。（判断）")
    doc.add_paragraph("答：正确")
    doc.add_paragraph("")

    doc.add_paragraph("第5题 一个圆的直径是 10 厘米，请计算它的周长和面积。（π 取 3.14）")
    doc.add_paragraph("答：")
    doc.add_paragraph("周长 = π × d = 3.14 × 10 = 31.4 厘米")
    doc.add_paragraph("")

    doc.add_paragraph("第6题 把一根绳子剪成两段，第一段长 3/5 米，第二段比第一段短 1/5 米，两根绳子一共长多少米？")
    doc.add_paragraph("答：")
    doc.add_paragraph("两根绳子一共长 3/5 + 1/5 = 4/5 米")

    doc.save(DOCX_PATH)
    print(f"✅ 学生答卷已生成：{DOCX_PATH}")


if __name__ == "__main__":
    asyncio.run(seed())
    gen_docx()
    print("\n演示路径：")
    print("  ① 学生(student01)登录 → 试卷批改 → 上传上面的 docx → 提交 → AI 批改")
    print("  ② 教师(teacher01)登录 → 试卷批改 → 待复核 → 通过发布 → 学生端看结果")
    print("\n若想让教师端「待复核列表」先有一条示例，起服务后跑：")
    print("  .venv/Scripts/python.exe scripts/demo_prepare_review.py")
